#!/usr/bin/env python3
"""Convert the user's monthly Excel layout into a Brújula portable JSON bundle.

Usage:
  python3 tools/excel_to_brujula.py input.xlsx output.json

The script reads only the workbook supplied by the user, never credentials or
bank connections. It understands the five four-column blocks used by the
Brújula Excel method: Servicios, Gastos Casa, Gastos supermercado,
Salidas/compras and Gastos auto.
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


BLOCKS = {
    "servicios": 1,
    "casa": 6,
    "supermercado": 11,
    "salidas": 16,
    "auto": 21,
}
MONTH_NAMES = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5,
    "junio": 6, "julio": 7, "agosto": 8, "septiembre": 9,
    "octubre": 10, "noviembre": 11, "diciembre": 12,
}


def parse_date(value, sheet_name: str) -> int:
    if isinstance(value, datetime):
        return int(value.timestamp() * 1000)
    if isinstance(value, date):
        return int(datetime(value.year, value.month, value.day).timestamp() * 1000)
    month = MONTH_NAMES.get(sheet_name.lower(), datetime.now().month)
    return int(datetime(datetime.now().year, month, 1).timestamp() * 1000)


def parse_amount(value):
    if isinstance(value, (int, float)):
        return int(round(value))
    if value is None:
        return None
    digits = re.sub(r"[^0-9-]", "", str(value))
    return int(digits) if digits else None


def fingerprint(sheet, category, received_at, merchant, amount):
    raw = f"excel|{sheet}|{category}|{received_at}|{merchant}|{amount}"
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def installment(text: str):
    match = re.search(r"(?:^|\s)(\d+)\s*[/ -]\s*(\d+)(?:\s|$)", text or "")
    return (int(match.group(1)), int(match.group(2))) if match else (0, 0)


def convert(source: Path):
    workbook = load_workbook(source, data_only=True, read_only=True)
    events = []
    commitments = []
    seen_commitments = set()
    for sheet_name in workbook.sheetnames:
        if sheet_name.lower() == "plantilla":
            continue
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(min_row=9, max_col=max(col + 3 for col in BLOCKS.values()), values_only=True)
        for row_values in rows:
            for category, first_column in BLOCKS.items():
                offset = first_column - 1
                when, _month, description, value = row_values[offset:offset + 4]
                merchant = str(description).strip() if description is not None else ""
                amount = parse_amount(value)
                if not merchant or not amount or amount <= 0:
                    continue
                received_at = parse_date(when, sheet_name)
                events.append({
                    "fingerprint": fingerprint(sheet_name, category, received_at, merchant, amount),
                    "received_at": received_at,
                    "source_package": "cl.brujula.app.excel",
                    "raw_text": f"EXCEL · {sheet_name} · {merchant} · {amount}",
                    "amount": amount,
                    "merchant": merchant,
                    "movement_type": "expense",
                    "category": category,
                    "payment_source": "excel",
                    "status": "requires_review",
                })
                current, total = installment(merchant)
                if total and merchant.lower() not in seen_commitments:
                    seen_commitments.add(merchant.lower())
                    commitments.append({
                        "id": hashlib.sha1(merchant.lower().encode("utf-8")).hexdigest(),
                        "name": merchant,
                        "amount": amount,
                        "current_installment": current,
                        "total_installments": total,
                        "category": category,
                        "recurring": True,
                    })
    return {
        "schema_version": 1,
        "generated_at": int(datetime.now().timestamp() * 1000),
        "events": events,
        "planning": {
            "schema_version": 1,
            "monthly_goal": 0,
            "incomes": [],
            "commitments": commitments,
        },
    }


def main():
    if len(sys.argv) != 3:
        raise SystemExit("Uso: python3 tools/excel_to_brujula.py entrada.xlsx salida.json")
    source, target = map(Path, sys.argv[1:])
    bundle = convert(source)
    target.write_text(json.dumps(bundle, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Convertidos {len(bundle['events'])} movimientos y {len(bundle['planning']['commitments'])} compromisos.")


if __name__ == "__main__":
    main()
