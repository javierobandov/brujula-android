#!/usr/bin/env python3
"""Smoke tests for the Excel import contract; never touches a user workbook."""

import json
import tempfile
from datetime import date
from pathlib import Path

from openpyxl import Workbook

from excel_to_brujula import convert


def build_fixture(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "enero"
    sheet.cell(9, 1, date(2026, 1, 5))
    sheet.cell(9, 3, "Luz 2/6")
    sheet.cell(9, 4, "$ 25.000")
    sheet.cell(9, 6, date(2026, 1, 6))
    sheet.cell(9, 8, "Arriendo")
    sheet.cell(9, 9, 450000)
    sheet.cell(10, 6, date(2026, 1, 7))
    sheet.cell(10, 8, "Arriendo")
    sheet.cell(10, 9, 450000)
    workbook.save(path)


def main() -> None:
    with tempfile.TemporaryDirectory() as directory:
        source = Path(directory) / "fixture.xlsx"
        build_fixture(source)
        bundle = convert(source)

    assert bundle["schema_version"] == 1
    assert len(bundle["events"]) == 3
    assert {event["category"] for event in bundle["events"]} == {"servicios", "casa"}
    assert all(event["status"] == "requires_review" for event in bundle["events"])
    commitments = bundle["planning"]["commitments"]
    assert len(commitments) == 1
    assert commitments[0]["current_installment"] == 2
    assert commitments[0]["total_installments"] == 6
    print(json.dumps({"events": len(bundle["events"]), "commitments": len(commitments)}))


if __name__ == "__main__":
    main()
